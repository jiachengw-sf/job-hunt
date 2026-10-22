"""The dashboard and APPLICATIONS.md: 0, 1 and 50 applications, a Markdown snapshot, escaping."""

from __future__ import annotations

import json
import os
import re
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest

from jobhunt.config import default_config_data
from jobhunt.errors import ExitCode
from jobhunt.models import Config
from jobhunt.store import connect, migrate
from jobhunt.tracker.applications import record_application
from jobhunt.tracker.dashboard import render_html, render_markdown
from jobhunt.tracker.sync import set_user_field
from jobhunt.tracker.timefmt import fmt_local, zone
from jobhunt.tracker.views import all_views, summary
from tests.conftest import REPO

NOW = "2026-10-22T04:42:00Z"  # Wed Oct 21, 9:42 PM in Los Angeles
LA = zone("America/Los_Angeles")
SNAPSHOT = REPO / "tests" / "fixtures" / "tracking" / "APPLICATIONS.md"
STATUSES = ["Applied", "OA", "Interviewing", "Rejected", "Offer", "Applied", "Withdrawn"]


def _cfg() -> Config:
    data = default_config_data()
    data["tracking"]["timezone"] = "America/Los_Angeles"
    return Config.model_validate(data)


def _store(tmp_path: Path, n: int, *, company=lambda i: f"Company {i:02d}"):
    c = connect(tmp_path / f"d{n}.db")
    migrate(c)
    for i in range(1, n + 1):
        ats = ["greenhouse", "lever", "workday", "ashby"][i % 4]
        url = (
            f"https://acme{i}.wd5.myworkdayjobs.com/External/job/X/Intern_R{i}"
            if ats == "workday"
            else f"https://example.com/{ats}/{i}"
        )
        c.execute(
            "INSERT INTO jobs (id, dedupe_key, ats, canonical_url, company, title, "
            "locations_json, work_mode, employment_type, first_seen, last_seen) "
            "VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (i, f"k{i}", ats, url, company(i), f"Software Engineer Intern {i}",
             json.dumps(["New York, NY"]), "hybrid", "Internship", NOW, NOW),
        )  # fmt: skip
        applied = (datetime(2026, 10, 21, 17, tzinfo=UTC) - timedelta(days=(i * 4) % 45)).strftime(
            "%Y-%m-%dT%H:%M:%SZ"
        )
        record_application(
            c, i, now=applied, applied_at=applied, approx=i % 5 == 0, method="manual",
            via="cli", tz=LA, paths=None, window_days=90, max_per_company=99,
        )  # fmt: skip
        st = STATUSES[i % len(STATUSES)]
        if st != "Applied":
            set_user_field(c, i, "status", st, applied.replace("T17", "T20"), source="user")
        if i in (2, 3):  # 3 is Rejected: its interview no longer needs attention
            c.execute(
                "INSERT INTO application_dates (app_id, kind, at, label) VALUES "
                "(?, 'interview', '2026-10-28T17:00:00Z', 'Onsite')",
                (i,),
            )
    return c


def _render(c):
    views = all_views(c, _cfg(), now=NOW, tz=LA)
    s = summary(views, now=NOW, tz=LA)
    gen = fmt_local(NOW, LA)
    return views, s, render_markdown(views, s, generated=gen), render_html(views, s, generated=gen)


def _embedded(html: str) -> dict:
    m = re.search(r'<script type="application/json" id="data">(.*?)</script>', html, re.S)
    assert m
    return json.loads(m.group(1))


def test_zero_applications(tmp_path):
    _, s, md, html = _render(_store(tmp_path, 0))
    assert s["total"] == 0 and s["interview_rate"] is None
    assert "No applications yet" in md and "Nothing right now." in md
    assert _embedded(html)["apps"] == []


def test_one_application(tmp_path):
    _, s, md, html = _render(_store(tmp_path, 1))
    assert s["total"] == 1 and "## OA (1)" in md
    data = _embedded(html)
    (app,) = data["apps"]
    assert app["company"] == "Company 01" and app["history"][-1]["to"] == "OA"
    assert "job_key" not in json.dumps(data) and "k1" not in json.dumps(data)


def test_fifty_applications(tmp_path):
    views, s, md, html = _render(_store(tmp_path, 50))
    assert s["total"] == 50 and len(_embedded(html)["apps"]) == 50
    assert sum(s["by_status"].values()) == 50
    assert md.count("| Company | Position |") == len(s["by_status"])
    rows = [ln for ln in md.splitlines() if ln.startswith("| Company ")]
    assert rows and all(ln.count("|") == 7 for ln in md.splitlines() if ln.startswith("| C"))
    assert (
        s["upcoming_interviews"]
        and "Interview: Company 02" in md
        and "Company 03 —" not in md.split("## ")[2]
    )
    # newest first inside each group
    group = [v for v in views if v["status"] == "Applied"]
    assert [v["applied_at"] for v in group] == sorted((v["applied_at"] for v in group),
                                                      reverse=True)  # fmt: skip


def test_markdown_snapshot(tmp_path):
    _, _, md, _ = _render(_store(tmp_path, 8))
    if os.environ.get("UPDATE_SNAPSHOTS"):
        SNAPSHOT.write_text(md, encoding="utf-8")
    assert md == SNAPSHOT.read_text(encoding="utf-8"), (
        "APPLICATIONS.md changed; rerun with UPDATE_SNAPSHOTS=1 if that's intended"
    )


def test_posting_text_cannot_inject_html(tmp_path):
    evil = '</script><img src=x onerror="alert(1)">Acme & <b>Co</b>'
    html = _render(_store(tmp_path, 1, company=lambda i: evil))[3]
    assert "</script><img" not in html and "<b>Co" not in html
    assert _embedded(html)["apps"][0]["company"] == evil  # shown as text by the page
    assert "innerHTML" not in html


def test_dashboard_command_and_sheet_edit_in_timeline(cli, persona_home):
    from openpyxl import load_workbook

    from tests.test_applications import WORKDAY, setup

    home = setup(cli, persona_home)
    assert (home / "dashboard.html").exists()  # `run` regenerates it, even with none yet
    (home / "dashboard.html").unlink()
    code, env, _ = cli("dashboard", "--dry-run")
    assert code == ExitCode.OK and not (home / "dashboard.html").exists()
    cli("applied", WORKDAY, "--at", "Oct 14 9:30pm")
    assert (home / "dashboard.html").exists() and (home / "APPLICATIONS.md").exists()
    cli("sync")
    xlsx = home / "tracker.xlsx"
    wb = load_workbook(xlsx)
    ws = wb["Applications"]
    header = [c.value for c in ws[1]]
    ws.cell(2, header.index("Status") + 1, "Rejected")
    wb.save(xlsx)
    cli("sync")
    code, env, _ = cli("dashboard", "--no-open")
    assert code == ExitCode.OK and env["data"]["applications"] == 1 and not env["data"]["opened"]
    app = _embedded((home / "dashboard.html").read_text(encoding="utf-8"))["apps"][0]
    assert app["status"] == "Rejected"
    assert app["history"][-1] | {"when": None} == {
        "when": None, "from": "Applied", "to": "Rejected", "who": "you, in the sheet",
        "note": None,
    }  # fmt: skip
    md = (home / "APPLICATIONS.md").read_text(encoding="utf-8")
    assert "## Rejected (1)" in md and "Harborlight Bank" in md


@pytest.mark.parametrize("n", [0, 1])
def test_views_written_by_refresh(tmp_path, n):
    from jobhunt.paths import Paths
    from jobhunt.tracker.views import refresh_views

    paths = Paths(home=tmp_path / "home")
    written = refresh_views(_store(tmp_path, n), paths, _cfg(), now=NOW)
    assert str(paths.dashboard) in written and paths.applications_md.exists()
