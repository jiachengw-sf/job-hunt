"""Regenerate the synthetic resume files (PDF and DOCX) from each persona's resume.md:

    uv run python tests/fixtures/resumes/make_resumes.py

The PDF writer is a minimal text-only PDF (Helvetica, one line per row), enough for pypdf to
extract; the DOCX uses python-docx. Every persona is synthetic.
"""

from __future__ import annotations

import re
from pathlib import Path

import docx

HERE = Path(__file__).resolve().parent
PERSONAS = HERE.parent.parent / "personas"


def plain_lines(md: str) -> list[str]:
    out = []
    for ln in md.splitlines():
        ln = re.sub(r"^#+\s*", "", ln)
        ln = ln.replace("**", "")
        out.append(ln.rstrip())
    return out


def _pdf_escape(s: str) -> str:
    s = s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    return s.encode("latin-1", "replace").decode("latin-1")


def write_pdf(lines: list[str], path: Path) -> None:
    rows_per_page = 52
    pages = [lines[i : i + rows_per_page] for i in range(0, len(lines), rows_per_page)] or [[]]
    objs: list[bytes] = []
    # 1 catalog, 2 pages, 3 font, then (content, page) per page
    kids = []
    for n, page in enumerate(pages):
        content_id = 4 + 2 * n
        page_id = content_id + 1
        kids.append(page_id)
        ops = ["BT", "/F1 10 Tf", "13 TL", "50 760 Td"]
        for ln in page:
            ops.append(
                f"({_pdf_escape(ln.replace('–', '-').replace('—', '-').replace('·', '|'))}) Tj T*"
            )
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        objs.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        objs.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>"
            % content_id
        )
    head = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids ["
        + b" ".join(b"%d 0 R" % k for k in kids)
        + b"] /Count %d >>" % len(kids),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    all_objs = head + objs
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(all_objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(all_objs) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(all_objs) + 1,
        xref,
    )
    path.write_bytes(bytes(out))


def write_docx(lines: list[str], path: Path) -> None:
    d = docx.Document()
    for ln in lines:
        if ln.startswith("- "):
            d.add_paragraph(ln[2:], style="List Bullet")
        else:
            d.add_paragraph(ln)
    d.core_properties.author = "jobhunt fixtures"
    d.save(str(path))


def main() -> None:
    for p in sorted(PERSONAS.iterdir()):
        md = p / "resume.md"
        if not md.exists():
            continue
        lines = plain_lines(md.read_text(encoding="utf-8"))
        write_pdf(lines, HERE / f"{p.name}.pdf")
        write_docx(lines, HERE / f"{p.name}.docx")
        print("wrote", p.name)


if __name__ == "__main__":
    main()
