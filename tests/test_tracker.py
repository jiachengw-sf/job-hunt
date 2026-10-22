"""Tracker: state machine, three-way sync, XLSX round trip (user edits survive), Sheets backend."""

from __future__ import annotations

import json
import sqlite3
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from jobhunt.errors import NeedsUser, ValidationFailed
from jobhunt.store.db import migrate
from jobhunt.tracker.backends import SheetsBackend, TrackerLocked, XlsxBackend
from jobhunt.tracker.states import machine
from jobhunt.tracker.store import agent_advance, mark
from jobhunt.tracker.sync import COLUMNS, ID_COL, sync

NOW = "2026-10-18T12:00:00Z"
TODAY = date(2026, 10, 18)


# --- state machine --------------------------------------------------------------------------


def test_agent_moves_forward_only_and_never_out_of_terminal():
    m = machine()
    assert m.agent_can_move("New", "Prefilled")
    assert not m.agent_can_move("Applied", "Prefilled")
    assert not m.agent_can_move("Rejected", "Interviewing")
    assert not m.agent_can_move("Skipped", "Applied")


def test_user_can_set_anything_but_skip_needs_a_reason():
    m = machine()
    assert m.user_move("new", None).name == "New"
    with pytest.raises(NeedsUser):
        m.user_move("Skipped", None)
    with pytest.raises(ValidationFailed):
        m.user_move("Skipped", "bored")
    assert m.user_move("Skipped", "pay").name == "Skipped"
    with pytest.raises(ValidationFailed):
        m.user_move("Hired!!", None)


# --- in-memory grid -------------------------------------------------------------------------


class FakeGrid:
    def __init__(self, rows: list[list[str]] | None = None) -> None:
        self.rows = [list(r) for r in rows or []]
        self.writes = 0

    def values(self):
        return [list(r) for r in self.rows]

    def set_cells(self, updates):
        for r, c, v in updates:
            while len(self.rows) < r:
                self.rows.append([])
            row = self.rows[r - 1]
            while len(row) < c:
                row.append("")
            row[c - 1] = v
            self.writes += 1

    def append_rows(self, rows):
        self.rows += [list(r) for r in rows]
        self.writes += len(rows)

    def replace_all(self, rows):
        self.rows = [list(r) for r in rows]


class FakeBackend:
    name, location = "fake", "memory"

    def __init__(self):
        self.tabs: dict[str, FakeGrid] = {}

    def tab(self, title, **_):
        return self.tabs.setdefault(title, FakeGrid())

    def save(self):
        pass


@pytest.fixture
def conn():
    c = sqlite3.connect(":memory:", isolation_level=None)
    c.row_factory = sqlite3.Row
    c.execute("PRAGMA foreign_keys = ON")
    migrate(c)
    for i, (company, title) in enumerate([("Acme", "Intern A"), ("Beta", "Intern B")], start=1):
        c.execute(
            "INSERT INTO jobs (id, dedupe_key, ats, company, title, locations_json, first_seen, "
            "last_seen) "
            "VALUES (?, ?, 'greenhouse', ?, ?, '[\"Chicago, IL\"]', ?, ?)",
            (i, f"k{i}", company, title, NOW, NOW),
        )
        c.execute(
            "INSERT INTO fits (job_id, gate_status, priority, band, provisional, computed_at) "
            "VALUES (?, 'passed', ?, 'good', 1, ?)",
            (i, 0.9 - i / 10, NOW),
        )
    yield c
    c.close()


def do_sync(conn, backend, dry_run=False):
    return sync(
        conn,
        backend,
        now=NOW,
        today=TODAY,
        window_days=90,
        include_gated_out=False,
        dry_run=dry_run,
    )


def col(name):
    return [c.header for c in COLUMNS].index(name)


def row_for(grid, job_id):
    header = grid.rows[0]
    idc = header.index(ID_COL)
    return next(r for r in grid.rows[1:] if r[idc] == str(job_id))


def cell(grid, job_id, header):
    r = row_for(grid, job_id)
    i = grid.rows[0].index(header)
    return r[i] if i < len(r) else ""


def set_cell(grid, job_id, header, value):
    r = row_for(grid, job_id)
    i = grid.rows[0].index(header)
    while len(r) <= i:
        r.append("")
    r[i] = value


def status(conn, job_id):
    return conn.execute("SELECT status FROM applications WHERE job_id=?", (job_id,)).fetchone()[0]


def test_first_sync_adds_rows_and_is_then_idempotent(conn):
    b = FakeBackend()
    rep = do_sync(conn, b)
    jobs = b.tabs["Jobs"]
    assert rep.rows_added == 2 and jobs.rows[0][: len(COLUMNS)] == [c.header for c in COLUMNS]
    assert cell(jobs, 1, "Status") == "New" and cell(jobs, 1, "Company") == "Acme"
    assert cell(jobs, 1, "Qualification band") == "good (provisional)"
    rep2 = do_sync(conn, b)
    assert rep2.rows_added == 0 and rep2.cells_pushed == 0 and not rep2.user_edits_applied


def test_sheet_edit_wins_and_is_recorded(conn):
    b = FakeBackend()
    do_sync(conn, b)
    set_cell(b.tabs["Jobs"], 1, "Status", "applied")
    set_cell(b.tabs["Jobs"], 1, "Notes", "Talked to Dana at the fair")
    rep = do_sync(conn, b)
    assert status(conn, 1) == "Applied"
    assert (
        conn.execute("SELECT notes FROM applications WHERE job_id=1").fetchone()[0]
        == "Talked to Dana at the fair"
    )
    ev = conn.execute(
        "SELECT details_json FROM events WHERE kind='status_changed' AND job_id=1"
    ).fetchone()
    assert json.loads(ev[0]) == {"from": "New", "to": "Applied"}
    assert len(rep.user_edits_applied) == 2


def test_mark_is_pushed_not_reverted(conn):
    b = FakeBackend()
    do_sync(conn, b)
    mark(conn, 2, "Interviewing", reason=None, applied_on="2026-10-01", today=TODAY, now=NOW)
    rep = do_sync(conn, b)
    assert status(conn, 2) == "Interviewing"
    assert cell(b.tabs["Jobs"], 2, "Status") == "Interviewing"
    assert cell(b.tabs["Jobs"], 2, "Applied on") == "2026-10-01"
    assert not rep.conflicts


def test_both_changed_sheet_wins_with_conflict(conn):
    b = FakeBackend()
    do_sync(conn, b)
    mark(conn, 1, "OA", reason=None, applied_on=None, today=TODAY, now=NOW)
    set_cell(b.tabs["Jobs"], 1, "Status", "Rejected")
    rep = do_sync(conn, b)
    assert status(conn, 1) == "Rejected" and rep.conflicts


def test_invalid_sheet_values_are_rejected_not_applied(conn):
    b = FakeBackend()
    do_sync(conn, b)
    set_cell(b.tabs["Jobs"], 1, "Status", "Hired!!")
    set_cell(b.tabs["Jobs"], 1, "Applied on", "last tuesday")
    rep = do_sync(conn, b)
    assert status(conn, 1) == "New"
    assert len(rep.rejected) == 2
    assert cell(b.tabs["Jobs"], 1, "Status") == "Hired!!"  # left for the user to fix


def test_agent_columns_are_overwritten_user_extras_kept(conn):
    b = FakeBackend()
    do_sync(conn, b)
    jobs = b.tabs["Jobs"]
    jobs.rows[0].append("My column")
    row_for(jobs, 1).extend([""] * (len(jobs.rows[0]) - len(row_for(jobs, 1)) - 1) + ["keep me"])
    set_cell(jobs, 1, "Company", "Typo Inc")
    do_sync(conn, b)
    assert cell(jobs, 1, "Company") == "Acme"
    assert cell(jobs, 1, "My column") == "keep me"


def test_dry_run_writes_nothing(conn):
    b = FakeBackend()
    rep = do_sync(conn, b, dry_run=True)
    assert rep.rows_added == 2
    assert b.tabs["Jobs"].rows == []
    assert conn.execute("SELECT count(*) FROM tracker_sync").fetchone()[0] == 0


def test_agent_advance(conn):
    assert agent_advance(conn, 1, "Prefilled", NOW)
    mark(conn, 1, "Skipped", reason="pay", applied_on=None, today=TODAY, now=NOW)
    assert not agent_advance(conn, 1, "Applied", NOW)
    assert status(conn, 1) == "Skipped"


# --- XLSX -----------------------------------------------------------------------------------


def test_xlsx_user_edits_survive_sync(conn, tmp_path: Path):
    """Required by the spec: edits in user-owned columns survive every sync."""
    path = tmp_path / "tracker.xlsx"
    do_sync(conn, XlsxBackend(path))
    wb = load_workbook(path)
    ws = wb["Jobs"]
    header = [c.value for c in ws[1]]
    rown = next(
        r
        for r in range(2, ws.max_row + 1)
        if str(ws.cell(r, header.index(ID_COL) + 1).value) == "1"
    )
    ws.cell(rown, header.index("Status") + 1, "Applied")
    ws.cell(rown, header.index("Notes") + 1, "Referred by a TA")
    ws.cell(rown, header.index("Follow-up on") + 1, date(2026, 10, 25))  # a real date cell
    ws.cell(1, len(header) + 1, "My column")
    ws.cell(rown, len(header) + 1, "mine")
    wb.save(path)

    conn.execute("UPDATE fits SET priority=0.99 WHERE job_id=1")  # agent data changes too
    do_sync(conn, XlsxBackend(path))
    do_sync(conn, XlsxBackend(path))  # and again: nothing is lost on repeat syncs

    ws = load_workbook(path)["Jobs"]
    header = [c.value for c in ws[1]]
    vals = {h: ws.cell(rown, i + 1).value for i, h in enumerate(header)}
    assert vals["Status"] == "Applied"
    assert vals["Notes"] == "Referred by a TA"
    assert str(vals["Follow-up on"])[:10] == "2026-10-25"
    assert vals["My column"] == "mine"
    assert vals["Priority"] == "0.99"
    app = conn.execute(
        "SELECT status, notes, follow_up_on FROM applications WHERE job_id=1"
    ).fetchone()
    assert tuple(app) == ("Applied", "Referred by a TA", "2026-10-25")
    assert {"Jobs", "Companies", "Runs"} <= set(load_workbook(path).sheetnames)


def test_xlsx_open_in_excel_is_skipped(tmp_path: Path):
    path = tmp_path / "tracker.xlsx"
    (tmp_path / "~$tracker.xlsx").write_text("lock", encoding="utf-8")
    with pytest.raises(TrackerLocked) as e:
        XlsxBackend(path)
    assert e.value.user_actions


# --- Google Sheets (fake gspread) -----------------------------------------------------------


class FakeWorksheet:
    def __init__(self, spreadsheet=None, sheet_id=0):
        self.grid = FakeGrid()
        self.col_count = 26
        self.id = sheet_id
        self.spreadsheet = spreadsheet

    def get_all_values(self):
        return self.grid.values()

    def batch_update(self, body, value_input_option=None):
        from gspread.utils import a1_to_rowcol

        self.grid.set_cells([(*a1_to_rowcol(u["range"]), u["values"][0][0]) for u in body])

    def append_rows(self, rows, value_input_option=None):
        self.grid.append_rows(rows)

    def clear(self):
        self.grid.rows = []

    def update(self, range_name, values, value_input_option=None):
        self.grid.rows = [list(r) for r in values]

    def add_cols(self, n):
        self.col_count += n


class FakeSpreadsheet:
    def __init__(self):
        self.sheets: dict[str, FakeWorksheet] = {}
        self.requests: list[dict] = []

    def batch_update(self, body):
        self.requests += body["requests"]

    def worksheet(self, title):
        if title not in self.sheets:
            raise type("WorksheetNotFound", (Exception,), {})(title)
        return self.sheets[title]

    def add_worksheet(self, title, rows, cols):
        self.sheets[title] = FakeWorksheet(self, len(self.sheets))
        return self.sheets[title]


class FakeClient:
    def __init__(self):
        self.sheet = FakeSpreadsheet()

    def open_by_key(self, key):
        return self.sheet


def test_sheets_backend_round_trip(conn):
    pytest.importorskip("gspread")
    client = FakeClient()
    backend = SheetsBackend("abc123", "unused.json", client=client)
    do_sync(conn, backend)
    ws = client.sheet.sheets["Jobs"]
    assert ws.col_count >= len(COLUMNS)
    set_cell(ws.grid, 2, "Status", "Withdrawn")
    do_sync(conn, backend)
    assert status(conn, 2) == "Withdrawn"


def _with_application(conn):
    from jobhunt.config import default_config_data
    from jobhunt.models import Config
    from jobhunt.tracker.applications import record_application

    record_application(
        conn, 1, now=NOW, applied_at=NOW, approx=False, method="manual", via="cli", tz=None,
        paths=None, window_days=90, max_per_company=2,
    )  # fmt: skip
    return Config.model_validate(default_config_data())


def test_sheets_applications_tab_is_formatted_once(conn):
    pytest.importorskip("gspread")
    cfg = _with_application(conn)
    client = FakeClient()
    backend = SheetsBackend("abc123", "unused.json", client=client)
    sync(conn, backend, now=NOW, today=TODAY, window_days=90, include_gated_out=False,
         dry_run=False, cfg=cfg)  # fmt: skip
    apps = client.sheet.sheets["Applications"]
    assert apps.grid.rows[0][0] == "Company" and apps.grid.rows[1][0] == "Acme"
    kinds = [next(iter(r)) for r in client.sheet.requests]
    for k in ("updateSheetProperties", "setBasicFilter", "setDataValidation",
              "addConditionalFormatRule", "updateDimensionProperties", "updateCells"):  # fmt: skip
        assert k in kinds, k
    rules = sum(k == "addConditionalFormatRule" for k in kinds)
    client.sheet.requests.clear()
    sync(conn, backend, now=NOW, today=TODAY, window_days=90, include_gated_out=False,
         dry_run=False, cfg=cfg)  # fmt: skip
    kinds = [next(iter(r)) for r in client.sheet.requests]
    assert "addConditionalFormatRule" not in kinds and rules  # not added twice
    assert kinds == ["updateCells"]  # links refreshed
