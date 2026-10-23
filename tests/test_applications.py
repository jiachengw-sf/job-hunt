"""Application records: manual, `mark`, sheet and chat recording; local times."""

from __future__ import annotations

import json
import sqlite3
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from jobhunt.errors import ExitCode, ValidationFailed
from jobhunt.tracker.timefmt import fmt_local, local_day, parse_when, zone
from tests.conftest import REPO

BASE = str(REPO / "tests" / "fixtures" / "sources" / "base")
WORKDAY = "https://harborlight.wd5.myworkdayjobs.com/en-US/careers/job/Charlotte-NC/Technology-Intern_R-10023"
LA = zone("America/Los_Angeles")


def db(home: Path) -> sqlite3.Connection:
    c = sqlite3.connect(home / "jobhunt.db")
    c.row_factory = sqlite3.Row
    return c


def setup(cli, persona_home, persona: str = "first-year-intern", **config) -> Path:
    home = persona_home(persona)
    cfg = {"tracking": {"timezone": "America/Los_Angeles"}} | config
    (home / "config.yaml").write_text(yaml.safe_dump(cfg), encoding="utf-8")
    code, env, _ = cli("run", "--fixtures", BASE, "--no-inspect", "--no-sync", "--no-notify")
    assert code == ExitCode.OK, env
    return home


def records(home: Path) -> list[sqlite3.Row]:
    with db(home) as c:
        return c.execute("SELECT * FROM application_records ORDER BY id").fetchall()


def job_id(home: Path, key_like: str) -> int:
    with db(home) as c:
        return c.execute(
            "SELECT id FROM jobs WHERE dedupe_key LIKE ? ORDER BY id", (key_like,)
        ).fetchone()[0]


# --- time ---------------------------------------------------------------------------------


def test_utc_round_trips_to_local_display_across_dst():
    # 2026-11-01 02:00 PDT -> 01:00 PST
    assert fmt_local("2026-10-22T04:42:00Z", LA) == "Wed, Oct 21, 2026 · 9:42 PM PDT"
    assert fmt_local("2026-11-02T05:42:00Z", LA) == "Sun, Nov 1, 2026 · 9:42 PM PST"
    assert fmt_local("2026-11-01T08:30:00Z", LA) == "Sun, Nov 1, 2026 · 1:30 AM PDT"
    assert fmt_local("2026-11-01T09:30:00Z", LA) == "Sun, Nov 1, 2026 · 1:30 AM PST"
    # what you type in local time comes back as the same local time on either side
    for local, stored in (
        ("2026-10-31 21:30", "2026-11-01T04:30:00Z"),
        ("2026-11-02 21:30", "2026-11-03T05:30:00Z"),
    ):
        utc, approx = parse_when(local, "2026-11-10T00:00:00Z", LA)
        assert (utc, approx) == (stored, False)
        assert fmt_local(utc, LA).endswith("9:30 PM " + ("PDT" if "10-31" in local else "PST"))
    assert local_day("2026-10-22T04:42:00Z", LA).isoformat() == "2026-10-21"
    assert fmt_local("2026-10-22T04:42:00Z", LA, approx=True) == (
        "Wed, Oct 21, 2026 (time approximate)"
    )


@pytest.mark.parametrize(
    ("text", "utc", "approx"),
    [
        ("now", "2026-10-22T04:42:00Z", False),
        ("today 3pm", "2026-10-21T22:00:00Z", False),
        ("this afternoon", "2026-10-21T22:00:00Z", False),
        ("yesterday 21:30", "2026-10-21T04:30:00Z", False),
        ("Oct 21 9:42pm", "2026-10-22T04:42:00Z", False),
        ("Oct 21, 2026 at 9:42 PM", "2026-10-22T04:42:00Z", False),
        ("2026-10-21 21:30", "2026-10-22T04:30:00Z", False),
        ("2026-10-21T21:30", "2026-10-22T04:30:00Z", False),
        ("2026-10-22T04:00:00Z", "2026-10-22T04:00:00Z", False),
        ("2 days ago", "2026-10-19T19:00:00Z", True),
        ("10/20", "2026-10-20T19:00:00Z", True),
        ("monday", "2026-10-19T19:00:00Z", True),
    ],
)
def test_natural_local_times(text, utc, approx):
    # "now" is Wed Oct 21, 9:42 PM in Los Angeles
    assert parse_when(text, "2026-10-22T04:42:00Z", LA) == (utc, approx)


@pytest.mark.parametrize("text", ["sometime", "25:00", "13pm", "tomorrow 3pm", "Oct 30 10am"])
def test_unreadable_or_future_times_are_refused(text):
    with pytest.raises(ValidationFailed):
        parse_when(text, "2026-10-22T04:42:00Z", LA)


def test_future_times_for_interviews():
    when = parse_when("Oct 30 10am", "2026-10-22T04:42:00Z", LA, allow_future=True)
    assert when == ("2026-10-30T17:00:00Z", False)


# --- recording ------------------------------------------------------------------------------


def test_manual_applied_for_a_workday_job(cli, persona_home):
    home = setup(cli, persona_home)
    code, env, _ = cli("applied", WORKDAY + "?utm_source=x", "--at", "yesterday 9:30pm",
                       "--login-email", "candidate@example.com")  # fmt: skip
    assert code == ExitCode.OK, env
    d = env["data"]
    assert d["created"] and d["company"] and d["status"] == "Applied"
    assert d["where_to_check"] == "https://harborlight.wd5.myworkdayjobs.com/careers/userHome"
    assert "Workday Candidate Home" in d["how_to_check"]
    assert d["applied"].endswith("9:30 PM PDT")
    (r,) = records(home)
    assert (r["method"], r["recorded_via"], r["status_check_method"]) == ("manual", "cli", "portal")
    assert r["login_email"] == "candidate@example.com" and r["resume_file"] == "resume.md"
    folder = home / "applications" / r["folder"]
    assert (folder / "job_description.md").exists() and (folder / "resume.md").exists()
    readme = (folder / "README.md").read_text(encoding="utf-8")
    assert "userHome" in readme and "Status timeline" in readme
    with db(home) as c:
        app = c.execute("SELECT status, applied_on FROM applications WHERE job_id=?",
                        (r["job_id"],)).fetchone()  # fmt: skip
    assert app["status"] == "Applied" and app["applied_on"]

    code, again, _ = cli("applied", str(r["job_id"]))  # re-run: no second record
    assert code == ExitCode.OK and not again["data"]["created"]
    assert again["data"]["app_id"] == r["id"] and len(records(home)) == 1


def test_applied_unknown_url_points_to_add(cli, persona_home):
    setup(cli, persona_home)
    code, env, _ = cli("applied", "https://careers.example.com/jobs/42")
    assert code == ExitCode.NOT_FOUND
    assert "jobhunt add https://careers.example.com/jobs/42" in env["hints"]["next_actions"][0]


def test_applied_dry_run_writes_nothing(cli, persona_home):
    home = setup(cli, persona_home)
    code, env, _ = cli("applied", WORKDAY, "--dry-run")
    assert code == ExitCode.OK and env["data"]["created"] and env["data"]["dry_run"]
    assert records(home) == [] and not (home / "applications").exists()


def test_chat_and_portal_override(cli, persona_home):
    home = setup(cli, persona_home)
    gh = job_id(home, "greenhouse:%")
    code, env, _ = cli("applied", str(gh), "--via", "chat", "--at", "yesterday 2pm", "--portal",
                       "https://careers.example.com/me", "--check-note", "Sign in with Gmail",
                       "--referral", "A. Friend", "--note", "applied from my phone")  # fmt: skip
    assert code == ExitCode.OK, env
    (r,) = records(home)
    assert r["recorded_via"] == "chat" and r["method"] == "referral"
    assert r["status_check_url"] == "https://careers.example.com/me"
    assert r["status_check_note"] == "Sign in with Gmail"
    with db(home) as c:
        app = c.execute("SELECT referral_contact, notes FROM applications WHERE job_id=?",
                        (gh,)).fetchone()  # fmt: skip
    assert tuple(app) == ("A. Friend", "applied from my phone")
    cli("applied", str(gh), "--via", "chat")
    assert len(records(home)) == 1


def test_mark_applied_records_once(cli, persona_home):
    home = setup(cli, persona_home)
    jid = job_id(home, "greenhouse:%")
    code, env, _ = cli("mark", str(jid), "Applied", "--applied-on", "2026-10-01")
    assert code == ExitCode.OK, env
    (r,) = records(home)
    assert r["applied_at_approx"] and r["applied_at"].startswith("2026-10-01")
    cli("mark", str(jid), "OA")
    cli("mark", str(jid), "Applied")
    assert len(records(home)) == 1
    code, env, _ = cli("mark", str(job_id(home, "lever:%")), "Skipped", "--reason", "pay")
    assert code == ExitCode.OK and len(records(home)) == 1  # skipping isn't applying


def _set_sheet_status(xlsx: Path, job: int, status: str) -> None:
    wb = load_workbook(xlsx)
    ws = wb["Jobs"]
    header = [c.value for c in ws[1]]
    id_c, st_c = header.index("jobhunt ID") + 1, header.index("Status") + 1
    for row in range(2, ws.max_row + 1):
        if str(ws.cell(row, id_c).value) == str(job):
            ws.cell(row, st_c, status)
    wb.save(xlsx)


def test_sheet_status_applied_records_once(cli, persona_home):
    home = setup(cli, persona_home)
    code, env, _ = cli("sync")
    assert code == ExitCode.OK, env
    jid = job_id(home, "greenhouse:%")
    _set_sheet_status(home / "tracker.xlsx", jid, "Applied")
    code, env, _ = cli("sync")
    assert code == ExitCode.OK and env["data"]["applications_recorded"] == 1, env
    (r,) = records(home)
    assert r["recorded_via"] == "sheet" and r["applied_at_approx"] == 1
    _set_sheet_status(home / "tracker.xlsx", jid, "Interviewing")
    code, env, _ = cli("sync")
    assert env["data"]["applications_recorded"] == 0 and len(records(home)) == 1
    with db(home) as c:
        hist = [
            (e["source_id"], json.loads(e["details_json"])["to"])
            for e in c.execute(
                "SELECT * FROM events WHERE job_id=? AND kind='status_changed' ORDER BY id", (jid,)
            )
        ]
    assert hist == [("tracker", "Applied"), ("tracker", "Interviewing")]


def test_volume_guard_warning_on_record(cli, persona_home):
    home = setup(cli, persona_home, volume_guard={"max_applications_per_company": 1})
    with db(home) as c:
        rows = c.execute(
            "SELECT company, group_concat(id) AS ids FROM jobs GROUP BY company "
            "HAVING count(*) >= 2"
        ).fetchall()
    if not rows:
        pytest.skip("fixtures have no company with two jobs")
    a, b = rows[0]["ids"].split(",")[:2]
    _, first, _ = cli("applied", a)
    assert not first["data"]["warnings"]
    _, second, _ = cli("applied", b)
    assert second["data"]["created"]
    assert any("volume guard" in w for w in second["data"]["warnings"])
    assert any("volume guard" in w for w in second["hints"]["user_actions"])


# --- views, status and updates ------------------------------------------------------------


def _store(tmp_path: Path):
    from jobhunt.store import connect, migrate

    c = connect(tmp_path / "unit.db")
    migrate(c)
    c.execute(
        "INSERT INTO jobs (id, dedupe_key, ats, company, title, first_seen, last_seen) "
        "VALUES (1, 'greenhouse:acme:1', 'greenhouse', 'Acme', 'Intern', ?, ?)",
        ("2026-09-01T00:00:00Z", "2026-09-01T00:00:00Z"),
    )
    return c


def _view_at(c, now: str, **cfg_over):
    from jobhunt.config import default_config_data
    from jobhunt.models import Config
    from jobhunt.tracker.views import app_view

    data = default_config_data()
    data["tracking"] |= {"timezone": "America/Los_Angeles"} | cfg_over
    cfg = Config.model_validate(data)
    rec = c.execute("SELECT * FROM application_records").fetchone()
    return app_view(c, rec, cfg, now=now, tz=LA)


def _record(c, applied_at: str) -> None:
    from jobhunt.tracker.applications import record_application

    record_application(
        c, 1, now=applied_at, applied_at=applied_at, approx=False, method="manual", via="cli",
        tz=LA, paths=None, window_days=90, max_per_company=2,
    )  # fmt: skip


def test_stale_and_follow_up_boundary_days(tmp_path):
    c = _store(tmp_path)
    _record(c, "2026-10-01T17:00:00Z")  # Thu Oct 1, 10 AM in Los Angeles

    def day(n: int) -> str:  # 4 PM in Los Angeles, n days after Oct 1
        return (datetime(2026, 10, 1, 23, tzinfo=UTC) + timedelta(days=n)).strftime(
            "%Y-%m-%dT%H:%M:%SZ"
        )

    v = _view_at(c, day(6))
    assert not v["followup_due"] and v["followup_due_on"] == "2026-10-08"
    v = _view_at(c, day(7))
    assert v["followup_due"] and "follow-up due" in v["attention"][0]
    assert not _view_at(c, day(20))["stale"]
    v = _view_at(c, day(21))
    assert v["stale"] and v["status"] == "Applied"  # flagged, never changed
    assert any("likely no response" in a for a in v["attention"])
    assert not _view_at(c, day(10), stale_after_days=11)["stale"]
    assert _view_at(c, day(11), stale_after_days=11)["stale"]
    # two follow-ups sent: no more are due
    for n in (1, 2):
        c.execute(
            "INSERT INTO drafts (job_id, kind, key, status, original, text, created_at, "
            "updated_at, sent_at) VALUES (1, 'followup', ?, 'sent', 'x', 'x', ?, ?, ?)",
            (f"followup-{n}", day(7), day(7), day(7)),
        )
    assert not _view_at(c, day(30))["followup_due"]
    # any update resets the clock; a status other than Applied is never stale
    from jobhunt.tracker.sync import set_user_field

    set_user_field(c, 1, "status", "OA", day(25), source="user")
    assert not _view_at(c, day(60))["stale"]


def test_follow_up_date_you_set_wins(tmp_path):
    from jobhunt.tracker.sync import set_user_field

    c = _store(tmp_path)
    _record(c, "2026-10-01T17:00:00Z")
    set_user_field(c, 1, "follow_up_on", "2026-10-04", "2026-10-02T00:00:00Z", source="user")
    assert not _view_at(c, "2026-10-04T06:00:00Z")["followup_due"]  # still Oct 3 in LA
    assert _view_at(c, "2026-10-04T23:00:00Z")["followup_due"]


def test_apps_show_status_update_and_history(cli, persona_home):
    home = setup(cli, persona_home)
    code, env, _ = cli("apps")
    assert code == ExitCode.OK and env["data"]["summary"]["total"] == 0
    assert "jobhunt applied" in env["hints"]["next_actions"][0]
    cli("applied", WORKDAY, "--at", "Oct 14 9:30pm")
    gh = job_id(home, "greenhouse:%")
    cli("applied", str(gh), "--at", "yesterday 10am")
    code, env, _ = cli("app", "status", f"job:{gh}", "OA", "--note", "OA invite",
                       "--at", "yesterday 11am")  # fmt: skip
    assert code == ExitCode.OK, env
    d = env["data"]
    assert d["change"]["from"] == "Applied" and d["change"]["to"] == "OA"
    assert [(h["from"], h["to"], h["who"]) for h in d["history"]] == [
        ("New", "Applied", "you"),
        ("Applied", "OA", "you"),
    ]
    assert d["history"][-1]["note"] == "OA invite"
    code, env, _ = cli("app", "status", f"job:{gh}", "Skipped")
    assert code == ExitCode.NEEDS_USER  # Skipped needs a reason
    code, env, _ = cli("app", "update", f"job:{gh}", "--interview", "Oct 30 10am", "--label",
                       "Phone screen", "--recruiter-name", "Sam Recruiter", "--next-step",
                       "Prep system design", "--login-email", "candidate@example.com")  # fmt: skip
    assert code == ExitCode.OK, env
    v = env["data"]
    assert v["upcoming_interviews"][0]["label"] == "Phone screen"
    assert v["next_step"] == "Prep system design" and v["recruiter"] == "Sam Recruiter"
    code, env, _ = cli("apps", "--status", "OA")
    assert [a["status"] for a in env["data"]["apps"]] == ["OA"]
    s = env["data"]["summary"]
    assert s["total"] == 2 and s["by_status"] == {"OA": 1, "Applied": 1}
    assert s["upcoming_interviews"] and s["interview_rate"] == 0
    code, env, _ = cli("app", "show", "99")
    assert code == ExitCode.NOT_FOUND


def test_human_views_are_readable(cli, persona_home, capsys):
    from jobhunt.cli.main import run

    setup(cli, persona_home)
    cli("applied", WORKDAY, "--at", "Oct 14 9:30pm")
    run(["app", "show", "1", "--human"])
    out = capsys.readouterr().out
    assert "Wed, Oct 14, 2026 · 9:30 PM PDT" in out and "Timeline:" in out
    assert "{" not in out and "job_key" not in out
    run(["apps", "--human"])
    out = capsys.readouterr().out
    assert out.splitlines()[0].startswith("1 application(s)") and "Harborlight" in out


def test_sheet_edits_win_over_agent_moves(cli, persona_home):
    from jobhunt.store import open_store
    from jobhunt.tracker.store import agent_advance

    home = setup(cli, persona_home)
    cli("sync")
    jid = job_id(home, "greenhouse:%")
    _set_sheet_status(home / "tracker.xlsx", jid, "Rejected")
    cli("sync")
    with open_store(home / "jobhunt.db") as c:
        assert not agent_advance(c, jid, "Applied", "2026-10-22T00:00:00Z")  # never backward
    _, env, _ = cli("app", "show", f"job:{jid}")
    assert env["data"]["status"] == "Rejected"
    assert [h["who"] for h in env["data"]["history"]] == ["you, in the sheet"]


def test_export_csv_and_xlsx(cli, persona_home):
    import csv

    home = setup(cli, persona_home)
    cli("applied", WORKDAY, "--at", "Oct 14 9:30pm")
    code, env, _ = cli("apps", "export")
    assert code == ExitCode.OK, env
    with (home / "applications.csv").open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert rows[0]["Company"] == "Harborlight Bank" and rows[0]["Status"] == "Applied"
    assert rows[0]["Applied"] == "Wed, Oct 14, 2026 · 9:30 PM PDT"
    assert "job_key" not in rows[0] and "Where to check" in rows[0]
    out = home / "x" / "apps.xlsx"
    code, env, _ = cli("apps", "export", "--xlsx", "--to", str(out))
    assert code == ExitCode.OK and env["data"]["rows"] == 1
    ws = load_workbook(out)["Applications"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert ws.data_validations.dataValidation
    code, env, _ = cli("apps", "export", "--dry-run", "--to", str(home / "none.csv"))
    assert not (home / "none.csv").exists()


def _app_tab(xlsx: Path):
    ws = load_workbook(xlsx)["Applications"]
    header = [c.value for c in ws[1]]
    return ws, header


def test_applications_tab_round_trip(cli, persona_home):
    home = setup(cli, persona_home)
    cli("applied", WORKDAY, "--at", "Oct 14 9:30pm")
    code, env, _ = cli("sync")
    assert code == ExitCode.OK, env
    xlsx = home / "tracker.xlsx"
    ws, header = _app_tab(xlsx)
    assert header[:6] == ["Company", "Position", "Location", "Applied", "Method", "Status"]
    row = {h: c.value for h, c in zip(header, ws[2], strict=False)}
    assert row["Company"] == "Harborlight Bank" and row["Status"] == "Applied"
    assert row["Applied"] == "Wed, Oct 14, 2026 · 9:30 PM PDT"
    assert row["Where to check"].endswith("/careers/userHome")
    link = ws.cell(2, header.index("Where to check") + 1)
    assert link.hyperlink and link.hyperlink.target == row["Where to check"]
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert any("Status" not in str(dv.sqref) for dv in ws.data_validations.dataValidation)
    assert len(ws.conditional_formatting) >= 2
    # your edits in this tab come back to the store and show in the Jobs tab
    wb = load_workbook(xlsx)
    ws = wb["Applications"]
    ws.cell(2, header.index("Status") + 1, "Interviewing")
    ws.cell(2, header.index("Next step") + 1, "Onsite prep")
    wb.save(xlsx)
    code, env, _ = cli("sync")
    assert code == ExitCode.OK and len(env["data"]["user_edits_applied"]) == 2, env
    _, show, _ = cli("app", "show", "1")
    assert show["data"]["status"] == "Interviewing" and show["data"]["next_step"] == "Onsite prep"
    assert show["data"]["history"][-1]["who"] == "you, in the sheet"
    jobs = load_workbook(xlsx)["Jobs"]
    jh = [c.value for c in jobs[1]]
    jrow = next(r for r in jobs.iter_rows(min_row=2, values_only=True)
                if str(r[jh.index("jobhunt ID")]) == str(show["data"]["job_id"]))  # fmt: skip
    assert jrow[jh.index("Status")] == "Interviewing"
    # and a `jobhunt app status` change is pushed, not reverted
    cli("app", "status", "1", "Offer")
    cli("sync")
    ws, header = _app_tab(xlsx)
    assert ws.cell(2, header.index("Status") + 1).value == "Offer"
    code, env, _ = cli("sync")
    assert not env["data"]["user_edits_applied"] and not env["data"]["conflicts"]


def test_first_application_lands_on_row_two_of_an_empty_tab(cli, persona_home):
    """The tab is created empty by an earlier sync; the first record still goes on row 2."""
    home = setup(cli, persona_home)
    cli("sync")
    ws, _ = _app_tab(home / "tracker.xlsx")
    assert ws.max_row == 1
    cli("applied", WORKDAY)
    cli("sync")
    ws, _ = _app_tab(home / "tracker.xlsx")
    assert ws.cell(2, 1).value == "Harborlight Bank" and ws.max_row == 2


def test_applied_without_a_posting_link_then_the_posting_merges(cli, persona_home):
    """A confirmation email names company, title and a req id, and no link. Recording it needs
    no URL; adding the posting later (its URL carries the req id) is the same job."""
    home = setup(cli, persona_home)
    args = ["--company", "Juniper Ridge Labs", "--title", "2027 Hardware Engineering Intern",
            "--req-id", "JR2027047", "--at", "yesterday 9am"]  # fmt: skip
    code, env, _ = cli("applied", *args)
    assert code == ExitCode.OK, env
    d = env["data"]
    assert d["created"] and d["company"] == "Juniper Ridge Labs" and d["status"] == "Applied"
    code, again, _ = cli("applied", *args)
    assert code == ExitCode.OK and not again["data"]["created"]
    assert again["data"]["job_id"] == d["job_id"] and len(records(home)) == 1

    _, apps, _ = cli("apps")
    (row,) = apps["data"]["apps"]
    at = datetime.fromisoformat(row["applied_at"].replace("Z", "+00:00"))  # UTC, ISO 8601
    assert row["applied_on"] == local_day(row["applied_at"], LA).isoformat() and at.tzinfo

    url = ("https://juniperridge.wd1.myworkdayjobs.com/en-US/External/job/Santa-Clara-CA/"
           "2027-Hardware-Engineering-Intern_JR2027047")  # fmt: skip
    code, env, _ = cli("add", url, "--company", "Juniper Ridge Labs", "--title",
                       "2027 Hardware Engineering Intern")  # fmt: skip
    assert code == ExitCode.OK, env
    assert env["data"]["job_id"] == d["job_id"] and not env["data"]["created"]
    with db(home) as c:
        key = c.execute("SELECT dedupe_key FROM jobs WHERE id=?", (d["job_id"],)).fetchone()[0]
    assert key.startswith("url:https://juniperridge.wd1.myworkdayjobs.com/")
    assert len(records(home)) == 1

    other = url.replace("juniperridge", "otherco")  # same req id, another company
    code, env, _ = cli("add", other, "--company", "Other Co", "--title", "Intern")
    assert env["data"]["job_id"] != d["job_id"]


def test_applied_needs_a_job_or_company_and_title(cli, persona_home):
    setup(cli, persona_home)
    code, env, _ = cli("applied", "--company", "Juniper Ridge Labs")
    assert code == ExitCode.VALIDATION and "--title" in env["error"]["message"]
